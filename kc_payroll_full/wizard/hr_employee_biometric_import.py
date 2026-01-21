# -*- coding: utf-8 -*-

import base64
import io
import logging
from openpyxl import load_workbook
from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)


class HrEmployeeBiometricImport(models.TransientModel):
    _name = "hr.employee.biometric.import"
    _description = "Importar IDs Biométricos desde Excel"

    file_data = fields.Binary("Archivo Excel", required=True)
    file_name = fields.Char("Nombre de archivo")

    def action_import(self):
        self.ensure_one()

        # Contadores para reporte
        updated_count = 0
        created_count = 0
        not_found_count = 0
        error_count = 0

        # Listas para detalles
        created_employees = []
        not_found_names = []
        error_details = []

        try:
            # 1) Abrir el Excel
            data = base64.b64decode(self.file_data)
            wb = load_workbook(filename=io.BytesIO(data), data_only=True)
            sheet = wb.active

            # 2) Parámetros de columnas según el formato del Excel:
            # Columna A (0): ID Empleac (ID Biométrico)
            # Columna B (1): Nombre
            IDX_ID_BIOMETRICO = 0  # Columna A
            IDX_NOMBRE = 1  # Columna B

            _logger.info("=== INICIANDO LECTURA DEL EXCEL ===")

            # 3) Procesar cada fila
            for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if idx == 1:
                    _logger.info("Encabezados detectados: %s", row)
                    continue  # Saltar encabezado

                # Obtener valores de las columnas
                biometric_id = row[IDX_ID_BIOMETRICO] if len(row) > IDX_ID_BIOMETRICO else None
                nombre_empleado = row[IDX_NOMBRE] if len(row) > IDX_NOMBRE else None

                _logger.info("Fila %d: ID Biométrico=%s, Nombre=%s", idx, biometric_id, nombre_empleado)

                # Validar que ambos campos tengan valor
                if not biometric_id or not nombre_empleado:
                    _logger.warning("Fila %d omitida: ID Biométrico=%s, Nombre=%s", idx, biometric_id, nombre_empleado)
                    continue

                # Convertir ID biométrico a entero
                try:
                    biometric_id_int = int(biometric_id) if isinstance(biometric_id, (int, float, str)) else None
                    if biometric_id_int is None:
                        raise ValueError("ID biométrico no es un número válido")
                except (ValueError, TypeError) as e:
                    _logger.error("Error en fila %d: ID biométrico inválido '%s' - %s", idx, biometric_id, str(e))
                    error_count += 1
                    error_details.append({
                        'fila': idx,
                        'nombre': nombre_empleado,
                        'id_biometrico': biometric_id,
                        'error': f'ID biométrico inválido: {str(e)}'
                    })
                    continue

                # Limpiar nombre (eliminar espacios extra)
                nombre_empleado = str(nombre_empleado).strip()

                if not nombre_empleado:
                    _logger.warning("Fila %d omitida: Nombre vacío", idx)
                    continue

                # Buscar empleado por nombre exacto
                employee = self.env["hr.employee"].search([
                    ("name", "=", nombre_empleado)
                ], limit=1)

                if employee:
                    # Empleado encontrado, actualizar ID biométrico
                    try:
                        employee.write({
                            'biometric_user_id': biometric_id_int
                        })
                        updated_count += 1
                        _logger.info("✓ Empleado actualizado: %s (ID: %d) -> ID Biométrico: %s",
                                     employee.name, employee.id, biometric_id_int)
                    except Exception as e:
                        _logger.error("Error actualizando empleado %s: %s", nombre_empleado, str(e))
                        error_count += 1
                        error_details.append({
                            'fila': idx,
                            'nombre': nombre_empleado,
                            'id_biometrico': biometric_id_int,
                            'error': f'Error al actualizar: {str(e)}'
                        })
                else:
                    # Empleado no encontrado, crear nuevo
                    try:
                        new_employee = self.env["hr.employee"].create({
                            'name': nombre_empleado,
                            'biometric_user_id': biometric_id_int,
                            'active': True,
                        })
                        created_count += 1
                        created_employees.append(nombre_empleado)
                        _logger.info("✓ Empleado creado: %s (ID: %d) con ID Biométrico: %s",
                                     new_employee.name, new_employee.id, biometric_id_int)
                    except Exception as e:
                        _logger.error("Error creando empleado %s: %s", nombre_empleado, str(e))
                        error_count += 1
                        error_details.append({
                            'fila': idx,
                            'nombre': nombre_empleado,
                            'id_biometrico': biometric_id_int,
                            'error': f'Error al crear: {str(e)}'
                        })

        except Exception as e:
            _logger.error("Error general en importación: %s", str(e))
            raise UserError(_('Error al procesar el archivo: %s') % str(e))

        # Construir mensaje de resumen
        message_parts = []
        
        if updated_count > 0:
            message_parts.append(f"✅ {updated_count} empleado(s) actualizado(s)")
        
        if created_count > 0:
            message_parts.append(f"🆕 {created_count} empleado(s) creado(s):")
            for name in created_employees:
                message_parts.append(f"   • {name}")
        
        if error_count > 0:
            message_parts.append(f"❌ {error_count} error(es) encontrado(s)")
            if error_details:
                message_parts.append("\nDetalles de errores:")
                for error in error_details[:10]:  # Mostrar solo los primeros 10
                    message_parts.append(f"   • Fila {error['fila']}: {error['nombre']} - {error['error']}")
                if len(error_details) > 10:
                    message_parts.append(f"   ... y {len(error_details) - 10} error(es) más")

        message = "\n".join(message_parts) if message_parts else "No se procesaron registros."

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': 'Importación Completada',
                'message': message,
                'type': 'success' if error_count == 0 else 'warning',
                'sticky': True,
            }
        }

