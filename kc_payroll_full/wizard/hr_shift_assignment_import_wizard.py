# -*- coding: utf-8 -*-

import base64
import io

from openpyxl import load_workbook
from odoo import api, fields, models, _
from odoo.exceptions import ValidationError


class HrShiftAssignmentImport(models.TransientModel):
    _name = 'hr.shift.assignment.import'
    _description = 'Importar asignaciones de turnos'

    file_data = fields.Binary(string='Archivo Excel', required=True)
    file_name = fields.Char(string='Nombre de archivo')
    apply_date = fields.Date(
        string='Fecha a aplicar',
        required=True,
        default=fields.Date.today,
        help='Fecha que se asignará a todos los registros importados'
    )
    rotation_id = fields.Many2one(
        'hr.shift.rotation',
        string='Rotación',
        help='Rotación desde donde se lanzó el importador'
    )

    def _normalize_text(self, value):
        if value is None:
            return ''
        text = str(value).strip().lower()
        return (text.replace('á', 'a')
                    .replace('é', 'e')
                    .replace('í', 'i')
                    .replace('ó', 'o')
                    .replace('ú', 'u'))

    def _parse_shift(self, value):
        normalized = self._normalize_text(value)
        if normalized in ['dia', 'day']:
            return 'dia'
        if normalized in ['noche', 'night']:
            return 'noche'
        return False

    def _get_header_index(self, header_row):
        headers = [self._normalize_text(h) for h in header_row]
        idx_map = {name: idx for idx, name in enumerate(headers) if name}
        employee_idx = idx_map.get('empleado') or idx_map.get('employee') or idx_map.get('nombre')
        shift_idx = idx_map.get('turno') or idx_map.get('shift')

        if employee_idx is None:
            for idx, name in enumerate(headers):
                if 'emplead' in name or 'employee' in name:
                    employee_idx = idx
                    break
        if shift_idx is None:
            for idx, name in enumerate(headers):
                if 'turno' in name or 'shift' in name:
                    shift_idx = idx
                    break

        return employee_idx, shift_idx

    def action_import(self):
        self.ensure_one()
        if not self.file_data:
            raise ValidationError(_('Debe cargar un archivo Excel.'))

        try:
            decoded = base64.b64decode(self.file_data)
            workbook = load_workbook(filename=io.BytesIO(decoded), data_only=True)
            sheet = workbook.active
        except Exception:
            raise ValidationError(_('No se pudo leer el archivo. Verifique que sea un Excel válido (.xlsx).'))

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValidationError(_('El archivo Excel está vacío.'))

        header_row = rows[0]
        employee_idx, shift_idx = self._get_header_index(header_row)
        start_row = 1

        if employee_idx is None or shift_idx is None:
            # Intentar tratar la primera fila como datos sin encabezados
            first_employee = header_row[0] if len(header_row) > 0 else ''
            first_shift = header_row[1] if len(header_row) > 1 else ''
            if first_employee and self._parse_shift(first_shift):
                employee_idx = 0
                shift_idx = 1
                start_row = 0
            else:
                raise ValidationError(_('El Excel debe incluir las columnas "empleado" y "turno".'))

        employee_model = self.env['hr.employee']
        assignment_model = self.env['hr.employee.shift.assignment']
        employees_to_add = set()

        created = 0
        updated = 0
        errors = []

        for row_num, row in enumerate(rows[start_row:], start=1 + start_row):
            employee_name = (row[employee_idx] if len(row) > employee_idx else '') or ''
            employee_name = str(employee_name).strip()
            shift_raw = row[shift_idx] if len(row) > shift_idx else ''
            shift_period = self._parse_shift(shift_raw)

            if not employee_name:
                errors.append(_('Fila %s: nombre de empleado vacío.') % row_num)
                continue
            if not shift_period:
                errors.append(_('Fila %s: turno inválido "%s". Use Día o Noche.') % (row_num, shift_raw))
                continue

            employees = employee_model.search([('name', 'ilike', employee_name)], limit=2)
            if not employees:
                errors.append(_('Fila %s: empleado no encontrado "%s".') % (row_num, employee_name))
                continue
            if len(employees) > 1:
                errors.append(_('Fila %s: nombre ambiguo "%s".') % (row_num, employee_name))
                continue

            employee = employees[0]
            if self.rotation_id:
                employees_to_add.add(employee.id)
            existing = assignment_model.search([
                ('employee_id', '=', employee.id),
                ('date', '=', self.apply_date),
            ], limit=1)

            assignment_model.create_or_update(
                employee.id,
                self.apply_date,
                shift_period,
                state='applied',
                reason=_('Importación de turnos'),
                rotation_id=self.rotation_id.id if self.rotation_id else False,
            )
            if existing:
                updated += 1
            else:
                created += 1

        if self.rotation_id and employees_to_add:
            self.rotation_id.write({
                'employee_ids': [(4, emp_id) for emp_id in employees_to_add]
            })

        message_lines = [
            _('Registros creados: %s') % created,
            _('Registros actualizados: %s') % updated,
        ]
        if errors:
            message_lines.append(_('Errores (%s):') % len(errors))
            message_lines.extend(errors[:20])
            if len(errors) > 20:
                message_lines.append(_('...y %s más.') % (len(errors) - 20))

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Importación de Turnos'),
                'message': '\n'.join(message_lines),
                'type': 'warning' if errors else 'success',
                'sticky': True if errors else False,
            }
        }
